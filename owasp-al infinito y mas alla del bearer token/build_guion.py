# -*- coding: utf-8 -*-
"""Genera guion.docx a partir de guion_part1.md, guion_part2.md y guion_part3.md."""
import re
import glob
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT

TOKEN = re.compile(
    r'(\*\*(?:(?!\*\*).)+?\*\*'          # **bold**
    r'|\[[^\]]+\]\([^)]+\)'              # [text](url)
    r'|(?<![*])\*[^*]+\*(?![*])'         # *italic*
    r'|`[^`]+`)'                         # `code`
)

def _set_props(run, bold=False, italic=False, mono=False):
    run.bold = bold
    run.italic = italic
    if mono:
        run.font.name = 'Consolas'
        run.font.color.rgb = RGBColor(0xA0, 0x30, 0x30)

def add_hyperlink(paragraph, url, text):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hl = OxmlElement('w:hyperlink')
    hl.set(qn('r:id'), r_id)
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    c = OxmlElement('w:color'); c.set(qn('w:val'), '0563C1'); rPr.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.append(u)
    r.append(rPr)
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    r.append(t)
    hl.append(r)
    paragraph._p.append(hl)

def write_inline(paragraph, text, bold=False, italic=False):
    pos = 0
    for m in TOKEN.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos:m.start()])
            _set_props(run, bold=bold, italic=italic)
        tok = m.group(0)
        if tok.startswith('**'):
            run = paragraph.add_run(tok[2:-2])
            _set_props(run, bold=True, italic=italic)
        elif tok.startswith('['):
            mm = re.match(r'\[([^\]]+)\]\(([^)]+)\)', tok)
            add_hyperlink(paragraph, mm.group(2), mm.group(1))
        elif tok.startswith('`'):
            run = paragraph.add_run(tok[1:-1])
            _set_props(run, bold=bold, italic=italic, mono=True)
        else:
            run = paragraph.add_run(tok[1:-1])
            _set_props(run, bold=bold, italic=True)
        pos = m.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        _set_props(run, bold=bold, italic=italic)

def split_row(line):
    line = line.strip()
    if line.startswith('|'): line = line[1:]
    if line.endswith('|'): line = line[:-1]
    return [c.strip() for c in line.split('|')]

def is_sep(cells):
    return all(re.fullmatch(r'[:\-\s]+', c) for c in cells if c) and any(cells)

def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1 + len(body), cols=len(header))
    t.style = 'Table Grid'
    for j, cell in enumerate(header):
        p = t.rows[0].cells[j].paragraphs[0]
        write_inline(p, cell, bold=True)
    for i, row in enumerate(body, start=1):
        for j in range(len(header)):
            cell_text = row[j] if j < len(row) else ''
            write_inline(t.rows[i].cells[j].paragraphs[0], cell_text)

def main():
    sources = sorted(glob.glob('guion_part*.md'))
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(0.9)

    first_h1 = True
    table_buf = []

    def flush_table():
        nonlocal table_buf
        if table_buf:
            rows = [split_row(l) for l in table_buf]
            rows = [r for r in rows if not is_sep(r)]
            if rows:
                add_table(doc, rows)
                doc.add_paragraph()
            table_buf = []

    for path in sources:
        with open(path, encoding='utf-8') as fh:
            lines = fh.read().splitlines()
        for raw in lines:
            line = raw.rstrip()
            if line.startswith('|'):
                table_buf.append(line)
                continue
            flush_table()
            if not line.strip() or line.strip() == '---':
                continue
            if line.startswith('# '):
                if first_h1:
                    doc.add_heading(line[2:].strip(), level=0)
                    first_h1 = False
                else:
                    doc.add_page_break()
                    doc.add_heading(line[2:].strip(), level=1)
            elif line.startswith('## '):
                doc.add_heading(line[3:].strip(), level=2)
            elif line.startswith('### '):
                doc.add_heading(line[4:].strip(), level=3)
            elif line.startswith('> '):
                p = doc.add_paragraph(style='Quote')
                write_inline(p, line[2:])
            elif line.startswith('- '):
                p = doc.add_paragraph(style='List Bullet')
                write_inline(p, line[2:])
            else:
                p = doc.add_paragraph()
                write_inline(p, line)
        flush_table()

    doc.save('guion.docx')
    # Verificación
    check = Document('guion.docx')
    words = sum(len(p.text.split()) for p in check.paragraphs)
    for tbl in check.tables:
        for row in tbl.rows:
            for cell in row.cells:
                words += len(cell.text.split())
    print(f"OK guion.docx | parrafos={len(check.paragraphs)} tablas={len(check.tables)} palabras~{words}")

if __name__ == '__main__':
    main()
