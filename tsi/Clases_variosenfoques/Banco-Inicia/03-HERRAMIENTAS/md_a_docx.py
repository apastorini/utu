# -*- coding: utf-8 -*-
"""
md_a_docx.py — Convierte los Markdown del kit Banco-Inicia a .docx.

Uso:
    python 03-HERRAMIENTAS/md_a_docx.py                  # convierte todo el kit
    python 03-HERRAMIENTAS/md_a_docx.py ruta\archivo.md  # un solo archivo

Los .docx se generan espejando la estructura en la carpeta 09_Docx (ej.:
    Banco-Inicia/README.md                     -> Banco-Inicia/09_Docx/README.docx
    Banco-Inicia/02-ENTREGABLES/A-GOBERNAR/x.md -> Banco-Inicia/09_Docx/02-ENTREGABLES/A-GOBERNAR/x.docx )

Requiere: python-docx  (pip install python-docx)
"""

import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
except ImportError:
    raise SystemExit("Falta python-docx. Ejecutá: pip install python-docx")

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parent
SALIDA = RAIZ / "09_Docx"
EXCLUIR = {"09_Docx"}

DARK = RGBColor(0x1F, 0x3B, 0x63)
GRAY = RGBColor(0x60, 0x60, 0x60)


def add_inline(par, text):
    tokens = re.split(r"(\*\*.*?\*\*|\*.*?\*|`[^`]*`)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            run = par.add_run(tok[2:-2])
            run.bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            run = par.add_run(tok[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            run = par.add_run(tok[1:-1])
            run.italic = True
        else:
            par.add_run(tok)


def parse_table(rows):
    data = []
    for line in rows:
        line = line.strip().strip("|")
        if re.match(r"^[\s:\-|]+$", line):
            continue
        data.append([c.strip() for c in line.split("|")])
    if data:
        ncols = max(len(r) for r in data)
        data = [r + [""] * (ncols - len(r)) for r in data]
    return data


def render_doc(md_path, doc):
    lines = md_path.read_text(encoding="utf-8").splitlines()
    i, n = 0, len(lines)
    in_code, code_buf = False, []

    def flush_code():
        if code_buf:
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Inches(0.4)
            for cl in code_buf:
                run = par.add_run(cl)
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                run.font.color.rgb = GRAY
                run.add_break()

    while i < n:
        line = lines[i].rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                in_code, code_buf = False, []
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        m = re.match(r"^# (.*)$", stripped)
        if m:
            t = doc.add_heading(m.group(1), level=1)
            for r in t.runs:
                r.font.color.rgb = DARK
                r.font.size = Pt(18)
            i += 1
            continue
        m = re.match(r"^## (.*)$", stripped)
        if m:
            h = doc.add_heading(m.group(1), level=2)
            for r in h.runs:
                r.font.color.rgb = DARK
            i += 1
            continue
        m = re.match(r"^### (.*)$", stripped)
        if m:
            h = doc.add_heading(m.group(1), level=3)
            for r in h.runs:
                r.font.color.rgb = DARK
            i += 1
            continue
        m = re.match(r"^#### (.*)$", stripped)
        if m:
            h = doc.add_heading(m.group(1), level=4)
            for r in h.runs:
                r.font.color.rgb = DARK
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < n:
            j = i
            while j < n and lines[j].strip().startswith("|"):
                j += 1
            data = parse_table(lines[i:j])
            if len(data) >= 1:
                table = doc.add_table(rows=len(data), cols=len(data[0]))
                table.style = "Light Grid Accent 1"
                for ri, row in enumerate(data):
                    for ci, cell in enumerate(row):
                        c = table.cell(ri, ci)
                        c.text = ""
                        p = c.paragraphs[0]
                        add_inline(p, cell.replace("←", "<-").replace("→", "->"))
                        if ri == 0:
                            for run in p.runs:
                                run.bold = True
            i = j
            continue

        m = re.match(r"^(\s*)[-*] (.*)$", line)
        if m:
            item = m.group(2)
            if item.strip().startswith("[ ]"):
                item = "\u2610 " + item.strip()[3:].strip()
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, item)
            i += 1
            continue

        m = re.match(r"^(\s*)(\d+)[.)]\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, m.group(3))
            i += 1
            continue

        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            run = p.add_run(stripped.lstrip("> ").strip())
            run.italic = True
            run.font.color.rgb = GRAY
            i += 1
            continue

        if re.match(r"^-{3,}$", stripped) or re.match(r"^\*{3,}$", stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            i += 1
            continue

        if stripped.startswith("<!--"):
            i += 1
            continue
        if not stripped:
            i += 1
            continue

        p = doc.add_paragraph()
        add_inline(p, stripped)
        i += 1

    flush_code()


def convert_one(md_path):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
    render_doc(md_path, doc)
    rel = md_path.relative_to(RAIZ)
    docx_path = SALIDA / rel.with_suffix(".docx")
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))
    print(f"OK  {rel} -> 09_Docx/{rel.with_suffix('.docx')}")


def main():
    args = sys.argv[1:]
    if args:
        target = Path(args[0]).resolve()
        convert_one(target)
        return
    archivos = []
    for md in RAIZ.rglob("*.md"):
        if EXCLUIR.isdisjoint(md.parts):
            archivos.append(md)
    archivos.sort()
    if not archivos:
        raise SystemExit("No hay .md para convertir")
    for md in archivos:
        convert_one(md)
    print(f"\nConvertidos {len(archivos)} documentos -> {SALIDA}")


if __name__ == "__main__":
    main()
