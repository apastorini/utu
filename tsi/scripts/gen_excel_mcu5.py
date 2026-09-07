# -*- coding: utf-8 -*-
"""Genera plantillas Excel (xlsx) de apoyo para MCU 5.0 y la bitácora."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = r"C:\utu\utu\tsi\tareas\plantilla\mcu5\excel"
os.makedirs(OUT, exist_ok=True)

HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
FUN_FILL = PatternFill("solid", fgColor="2E74B5")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=11)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)


def estilo_hoja(ws, widths, freeze="A4"):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.freeze_panes = freeze


def write_rows(ws, rows, start=2):
    for r_i, row in enumerate(rows, start):
        for c_i, v in enumerate(row, 1):
            c = ws.cell(row=r_i, column=c_i, value=v)
            c.alignment = WRAP
            c.border = THIN
        fun = row[0]
        if fun in ("Gobernar", "Identificar", "Proteger", "Detectar", "Responder", "Recuperar"):
            for c_i in range(1, len(row) + 1):
                ws.cell(row=r_i, column=c_i).fill = FUN_FILL
                ws.cell(row=r_i, column=c_i).font = Font(bold=True, color="FFFFFF")


# ---------------------------------------------------------------- CONTROLES MCU 5.0
CONTROLS = [
    ("Función MCU 5.0", "Resultado/Control esperado", "Aplica (Sí/No/N.A.)", "Justificación si N.A.",
     "Evidencia necesaria", "Cómo se demuestra en la auditoría"),
]

gv = [
    ("Gobernar", "Política de ciberseguridad aprobada y comunicada", "", "", "Política firmada", "Presentar la política aprobada con fecha y firma de la Dirección"),
    ("Gobernar", "Estrategia de ciberseguridad alineada al negocio", "", "", "Estrategia/plan anual", "Mostrar el plan anual y su presentación a la Dirección"),
    ("Gobernar", "Roles y responsabilidades claros (RSI, comité, dueños)", "", "", "Matriz RACI / organigrama", "Desplegar la matriz RACI con responsables por proceso"),
    ("Gobernar", "Programa de gestión de riesgo institucional", "", "", "Análisis de riesgos", "Presentar matriz de riesgos con tratamiento y riesgo residual aceptado"),
    ("Gobernar", "Gestión de cadena de suministro y terceros", "", "", "Análisis de dependencias", "Listar servidores/proveedores y su evaluación"),
    ("Gobernar", "Marco de cumplimiento (MCU5, BCU, ISO, COBIT, URCDP)", "", "", "SoA + brecha MCU", "Mostrar la SoA y la brecha por función con perfil Avanzado"),
    ("Gobernar", "Presupuesto y recursos de seguridad", "", "", "Plan de inversión", "Presentar carta con dotación de recursos"),
    ("Gobernar", "Revisión periódica de gobierno por la Dirección", "", "", "Actas de comité", "Mostrar actas de reuniones de seguimiento"),
]
gv = [tuple(("Gobernar",) + t if False else (t[0],) + t[1:]) if False else t for t in gv]

ident = [
    ("Identificar", "Inventario de activos de información", "", "", "Registro de activos (Excel MCU)", "Desplegar el Excel de activos completo y actualizado"),
    ("Identificar", "Propietarios de activos y clasificación", "", "", "Registro de activos", "Mostrar que cada activo tiene dueño y clasificación"),
    ("Identificar", "Análisis de impacto en el negocio (BIA)", "", "", "Matriz de criticidad", "Presentar procesos críticos y activos que los soportan"),
    ("Identificar", "Evaluación de riesgos de ciberseguridad", "", "", "Análisis de riesgos", "Presentar matriz 5x5 con tratamiento del riesgo"),
    ("Identificar", "Comunicación y aceptación de riesgos", "", "", "Acta de aceptación", "Mostrar riesgo residual aceptado y firmado"),
    ("Identificar", "Inventario de vulnerabilidades", "", "", "Registro de vulnerabilidades", "Desplegar el registro con CVSS y estado"),
    ("Identificar", "Mapa de datos personales (URCDP)", "", "", "Registro de bases de datos", "Listar bases con datos personales y su tratamiento"),
    ("Identificar", "Identificación de dependencias externas", "", "", "Diagrama de arquitectura", "Mostrar servicios externos/integrados en el diagrama"),
]
ident = [t for t in ident]

proteger = [
    ("Proteger", "Gestión de identidades y credenciales", "", "", "Configuración IAM", "Demostrar alta/baja de usuarios y políticas de credenciales"),
    ("Proteger", "Autenticación multifactor (TOTP/WebAuthn/Hello)", "", "", "Config. + capturas", "Registrar login con 2 factores en vivo"),
    ("Proteger", "Revisión de privilegios (mínimo privilegio)", "", "", "Matriz de accesos", "Mostrar revisión periódica de privilegios"),
    ("Proteger", "Protección de datos (cifrado en reposo y en tránsito)", "", "", "Cifrado de BD y TLS", "Demostrar cifrado de bóveda/BD y TLS 1.2+"),
    ("Proteger", "Respaldos y prueba de restauración", "", "", "Registro de backup", "Restaurar un respaldo en vivo ante la auditoría"),
    ("Proteger", "Capacitación y concientización del personal", "", "", "Plan de capacitación", "Presentar plan y registro de asistencia"),
    ("Proteger", "Configuración segura de sistemas", "", "", "Config. endurecida", "Mostrar hardening por checklist"),
    ("Proteger", "Mantenimiento y parcheo", "", "", "Historial de parches", "Mostrar aplicaciones de parches en el período"),
    ("Proteger", "Protección de endpoints (HIDS/antivirus/EDR)", "", "", "Agentes activos", "Mostrar agentes HIDS conectados al SIEM"),
    ("Proteger", "Protección de la red (firewalls, segmentación, IDS/IPS)", "", "", "Diagrama + reglas", "Presentar reglas de firewall y segmentación en vivo"),
    ("Proteger", "Gestión de vulnerabilidades técnicas", "", "", "Escaneos + cierre", "Mostrar escaneos y resolución con plazos"),
    ("Proteger", "Seguridad en el ciclo de desarrollo (SAST)", "", "", "Informe SAST", "Mostrar informe de semgrep/bandit sin hallazgos críticos"),
]
proteger = [t for t in proteger]

detectar = [
    ("Detectar", "Monitoreo continuo de la red (NIDS)", "", "", "Reglas Suricata/Zeek", "Mostrar eventos de red capturados recientemente"),
    ("Detectar", "Monitoreo de hosts (HIDS/FIM)", "", "", "Alertas Wazuh", "Mostrar monitoreo de integridad de archivos"),
    ("Detectar", "Correlación de eventos (SIEM)", "", "", "Dashboards SIEM", "Presentar dashboard con eventos correlacionados"),
    ("Detectar", "Firmas y reglas de detección actualizadas", "", "", "Versión de firmas", "Mostrar fecha de actualización de firmas"),
    ("Detectar", "Pruebas de detección (ejercicios púrpura/red team)", "", "", "Evidencia de simulación", "Presentar al menos un ataque simulado detectado"),
    ("Detectar", "Detección de anomalías de comportamiento", "", "", "Regla de anomalía", "Mostrar una regla de comportamiento y su alerta"),
    ("Detectar", "Alerta y notificación oportuna", "", "", "Canales de alerta", "Demostrar una alerta por mail/webhook/llamada"),
    ("Detectar", "Monitoreo de los controles de protección", "", "", "Estado de controles", "Mostrar salud de firewalls/agentes en el dashboard"),
]
detectar = [t for t in detectar]

responder = [
    ("Responder", "Plan de respuesta a incidentes", "", "", "Procedimiento de incidentes", "Exponer el plan y sus fases"),
    ("Responder", "Roles y comunicación en respuesta", "", "", "Matriz de comunicación", "Mostrar responsables de comunicación y escalamiento"),
    ("Responder", "Análisis y contención de incidentes", "", "", "Registros de incidentes", "Presentar un incidente completo (detección a cierre)"),
    ("Responder", "Erradicación y recuperación", "", "", "Incidente resuelto", "Mostrar evidencia de la mitigación"),
    ("Responder", "Notificación a autoridades (BCU, URCDP, CERTuy)", "", "", "Borradores de notificación", "Presentar notificación simulada completa"),
    ("Responder", "Lecciones aprendidas y mejora", "", "", "Informe de lecciones", "Mostrar acciones correctivas aplicadas"),
]
responder = [t for t in responder]

recuperar = [
    ("Recuperar", "Plan de recuperación (BCP/DRP)", "", "", "Plan de continuidad", "Exponer el BCP/DRP con RTO/RPO"),
    ("Recuperar", "Comunicación de crisis", "", "", "Procedimiento de crisis", "Presentar canales y responsables de crisis"),
    ("Recuperar", "Restauración de respaldos validada", "", "", "Prueba de restauración", "Ejecutar/evidenciar una restauración real"),
    ("Recuperar", "Reanudación de operaciones", "", "", "Plan de retorno", "Mostrar secuencia de reanudación de servicios"),
    ("Recuperar", "Mejora continua post-recuperación", "", "", "Acciones de mejora", "Presentar mejoras aplicadas tras las simulaciones"),
]
recuperar = [t for t in recuperar]

wb = Workbook()
ws = wb.active
ws.title = "Controles MCU5"
ws.append(CONTROLS[0])
rows = gv + ident + proteger + detectar + responder + recuperar
write_rows(ws, rows)
estilo_hoja(ws, [12, 52, 14, 30, 34, 50])
path = os.path.join(OUT, "01-controles-mcu5-perfil-avanzado.xlsx")
wb.save(path)
print("OK", path)

# ---------------------------------------------------------------- ACTIVOS
wb = Workbook()
ws = wb.active
ws.title = "Activos"
ws.append([
    "ID activo", "Nombre del activo", "Tipo (HW/SW/Dato/Servicio)", "Descripción",
    "Ubicación", "Dueño", "Clasificación", "Criticidad", "Proceso de negocio",
    "Dato personal (S/N)", "Función MCU vinculada", "Comentarios",
])
activos = [
    ("A01", "Servidor de control central", "HW/SW", "VM con el control central (API + dashboard)", "VLAN Servidores", "RSI", "Confidencial", "Alta", "Gestión de secretos", "N", "Proteger/Detectar", ""),
    ("A02", "PostgreSQL (eventos)", "Dato", "Base de eventos de auditoría", "Servidor control", "RSI", "Confidencial", "Alta", "Auditoría", "S", "Detectar", "Retención 90+ días"),
    ("A03", "Bóveda en cliente", "Dato", "Bóveda cifrada de credenciales", "Estación usuario", "Usuario", "Secreto", "Muy alta", "Acceso a sistemas", "S", "Proteger", "Solo cliente"),
    ("A04", "Servidor SIEM", "HW/SW", "Wazuh / Elastic (correlación)", "VLAN Seguridad", "RSI", "Confidencial", "Alta", "Monitoreo", "N", "Detectar", ""),
    ("A05", "Servidor de correo", "HW/SW", "Postfix/Dovecot (SMTP)", "VLAN Servicios", "RSI", "Interno", "Media", "Notificaciones", "S", "Responder", ""),
    ("A06", "Firewall de perímetro", "HW/SW", "OPNsense con VLANs y reglas", "Perímetro", "RSI", "Confidencial", "Alta", "Red", "N", "Proteger", ""),
    ("A07", "NIDS", "HW/SW", "Suricata/Zeek (mirror de tráfico)", "VLAN Seguridad", "RSI", "Confidencial", "Alta", "Detección", "N", "Detectar", ""),
    ("A08", "Base de datos de clientes", "Dato", "Datos personales de usuarios", "Servidor control", "RSI", "Secreto", "Muy alta", "Auditoría", "S", "Identificar/Proteger", "Registro URCDP"),
    ("A09", "Llave maestra / certificados", "Dato", "Material criptográfico", "HSM/cofre local", "RSI", "Secreto", "Muy alta", "Criptografía", "N", "Proteger", "Rotación definida"),
    ("A10", "Agentes HIDS (clientes)", "SW", "Agentes Wazuh en estaciones", "Estaciones", "RSI", "Interno", "Media", "Monitoreo de hosts", "N", "Detectar", ""),
]
write_rows(ws, activos)
estilo_hoja(ws, [9, 24, 22, 34, 18, 14, 16, 12, 20, 16, 24, 24])
path = os.path.join(OUT, "02-registro-activos-mcu5.xlsx")
wb.save(path)
print("OK", path)

# ---------------------------------------------------------------- RACI
wb = Workbook()
ws = wb.active
ws.title = "Matriz RACI"
ws.append(["Proceso/Actividad de seguridad", "R (Responsable)", "A (Accountable)", "C (Consultado)", "I (Informado)", "Documentos de referencia", "Evidencia"])
raci = [
    ("Gestión de incidentes", "Analista/SOC", "RSI", "Dueños de activos, TI", "Dirección", "Procedimiento de incidentes", "Registros de incidentes"),
    ("Gestión de accesos", "Admin IAM", "RSI", "Jefes de unidad", "Dueños de activos", "Política de accesos", "Matriz de accesos y revisión"),
    ("Evaluación de riesgos", "RSI", "Dirección", "Dueños de procesos", "Auditoría", "Metodología ISO 31000", "Matriz de riesgos"),
    ("Gestión de vulnerabilidades", "Equipo TI", "RSI", "Dueños de activos", "Dirección", "Procedimiento de parcheo", "Registro con CVSS/SLA"),
    ("Capacitación", "RRHH + RSI", "RSI", "Jefaturas", "Dirección", "Plan anual", "Registro de asistencia"),
    ("Continuidad (BCP/DRP)", "Tecnología", "RSI", "Dirección", "Jefaturas", "BIA + BCP", "Prueba de restauración"),
    ("Notificación a autoridades", "RSI", "Dirección", "Legal", "Dirección", "Procedimiento de notificación", "Borradores + constancias"),
]
write_rows(ws, raci)
estilo_hoja(ws, [34, 16, 16, 18, 14, 30, 34])
path = os.path.join(OUT, "03-matriz-raci-mcu5.xlsx")
wb.save(path)
print("OK", path)

# ---------------------------------------------------------------- BITÁCORA
wb = Workbook()
ws = wb.active
ws.title = "Bitácora"
ws.append(["Fecha", "Hora (UTC)", "Fase/Actividad", "Responsable", "Tarea realizada", "Herramienta/Comando", "Resultado / Evidencia", "Incidencia (S/N)", "Observaciones"])
bitacora = [
    ("", "", "Blue Team", "", "", "", "", "", ""),
    ("", "", "Blue Team", "", "", "", "", "", ""),
    ("", "", "Red Team", "", "", "", "", "", ""),
    ("", "", "Red Team", "", "", "", "", "", ""),
]
write_rows(ws, bitacora)
estilo_hoja(ws, [12, 12, 14, 16, 40, 28, 36, 14, 30])
path = os.path.join(OUT, "04-bitacora-planilla.xlsx")
wb.save(path)
print("OK", path)

print("Excell generados en:", OUT)