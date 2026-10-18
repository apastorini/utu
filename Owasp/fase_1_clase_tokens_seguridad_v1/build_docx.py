# -*- coding: utf-8 -*-
"""Convierte un .md (con encabezados, tablas, listas, citas y bloques de codigo) a .docx.
Uso: python build_docx.py <entrada.md> <salida.docx>
"""
import os
import re
import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT

IMG = re.compile(r'^!\[([^\]]*)\]\(([^)]+)\)\s*$')
OLIST = re.compile(r'^\d+\.\s+(.*)$')

TOKEN = re.compile(
    r'(\*\*(?:(?!\*\*).)+?\*\*'
    r'|\[[^\]]+\]\([^)]+\)'
    r'|(?<![*])\*[^*]+\*(?![*}])'
    r'|`[^`]+`)'
)

def _props(run, bold=False, italic=False, mono=False, size=None):
    run.bold = bold
    run.italic = italic
    if mono:
        run.font.name = 'Consolas'
        run.font.color.rgb = RGBColor(0x1F, 0x3B, 0x66)
    if size:
        run.font.size = Pt(size)

def add_hyperlink(paragraph, url, text):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hl = OxmlElement('w:hyperlink')
    hl.set(qn('r:id'), r_id)
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    c = OxmlElement('w:color'); c.set(qn('w:val'), '0563C1'); rPr.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.append(u)
    r.append(rPr)
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text
    r.append(t); hl.append(r); paragraph._p.append(hl)

def write_inline(paragraph, text, bold=False, italic=False):
    pos = 0
    for m in TOKEN.finditer(text):
        if m.start() > pos:
            _props(paragraph.add_run(text[pos:m.start()]), bold=bold, italic=italic)
        tok = m.group(0)
        if tok.startswith('**'):
            write_inline(paragraph, tok[2:-2], bold=True, italic=italic)
        elif tok.startswith('['):
            mm = re.match(r'\[([^\]]+)\]\(([^)]+)\)', tok)
            label, url = mm.group(1), mm.group(2)
            if url.startswith('#'):          # ancla interna: no es URL externa
                _props(paragraph.add_run(label), bold=bold, italic=italic)
            else:
                add_hyperlink(paragraph, url, label)
        elif tok.startswith('`'):
            _props(paragraph.add_run(tok[1:-1]), bold=bold, italic=italic, mono=True)
        else:
            write_inline(paragraph, tok[1:-1], bold=bold, italic=True)
        pos = m.end()
    if pos < len(text):
        _props(paragraph.add_run(text[pos:]), bold=bold, italic=italic)

def split_row(line):
    line = line.strip()
    if line.startswith('|'): line = line[1:]
    if line.endswith('|'): line = line[:-1]
    return [c.strip() for c in line.split('|')]

def is_sep(cells):
    return bool(cells) and all(re.fullmatch(r'[:\-\s]*', c) for c in cells)

def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1 + len(body), cols=len(header))
    t.style = 'Table Grid'
    for j, cell in enumerate(header):
        write_inline(t.rows[0].cells[j].paragraphs[0], cell, bold=True)
    for i, row in enumerate(body, start=1):
        for j in range(len(header)):
            write_inline(t.rows[i].cells[j].paragraphs[0], row[j] if j < len(row) else '')

def add_code(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.left_indent = Inches(0.15)
    for i, ln in enumerate(lines):
        run = p.add_run(ln)
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x1F, 0x3B, 0x66)
        if i < len(lines) - 1:
            run.add_break()

def add_heading_md(doc, text, level):
    p = doc.add_heading('', level=level)
    write_inline(p, text)
    return p

def add_image(doc, src_dir, alt, path, width=6.3):
    full = path if os.path.isabs(path) else os.path.join(src_dir, path)
    if not os.path.exists(full):
        write_inline(doc.add_paragraph(), f"[imagen no encontrada: {path}]")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(full, width=Inches(width))
    if alt:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(alt)
        r.italic = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(0x49, 0x50, 0x57)

def convert(src, dst):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Calibri'; st.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.85)

    lines = open(src, encoding='utf-8').read().splitlines()
    src_dir = os.path.dirname(os.path.abspath(src))
    table_buf, code_buf = [], []
    in_code = False
    first_h1 = True

    def flush_table():
        nonlocal table_buf
        if table_buf:
            rows = [split_row(l) for l in table_buf]
            rows = [r for r in rows if not is_sep(r)]
            if rows:
                add_table(doc, rows); doc.add_paragraph()
            table_buf = []

    def flush_code():
        nonlocal code_buf
        if code_buf:
            add_code(doc, code_buf); doc.add_paragraph()
            code_buf = []

    for raw in lines:
        line = raw.rstrip()
        if line.strip().startswith('```'):
            flush_table()
            if in_code:
                flush_code()
            in_code = not in_code
            continue
        if in_code:
            code_buf.append(raw)
            continue
        if line.startswith('|'):
            table_buf.append(line); continue
        flush_table()
        mimg = IMG.match(line.strip())
        if mimg:
            add_image(doc, src_dir, mimg.group(1), mimg.group(2))
            continue
        if not line.strip() or line.strip() == '---':
            continue
        if line.startswith('# '):
            if first_h1:
                add_heading_md(doc, line[2:].strip(), 0); first_h1 = False
            else:
                doc.add_page_break(); add_heading_md(doc, line[2:].strip(), 1)
        elif line.startswith('## '):
            add_heading_md(doc, line[3:].strip(), 2)
        elif line.startswith('### '):
            add_heading_md(doc, line[4:].strip(), 3)
        elif line.startswith('#### '):
            add_heading_md(doc, line[5:].strip(), 4)
        elif line.startswith('##### '):
            add_heading_md(doc, line[6:].strip(), 5)
        elif line.startswith('> '):
            write_inline(doc.add_paragraph(style='Quote'), line[2:])
        elif line.startswith('- '):
            write_inline(doc.add_paragraph(style='List Bullet'), line[2:])
        elif OLIST.match(line):
            write_inline(doc.add_paragraph(style='List Number'), OLIST.match(line).group(1))
        else:
            write_inline(doc.add_paragraph(), line)
    flush_table(); flush_code()
    doc.save(dst)

    c = Document(dst)
    words = sum(len(p.text.split()) for p in c.paragraphs)
    for t in c.tables:
        for row in t.rows:
            for cell in row.cells:
                words += len(cell.text.split())
    links = sum(1 for r in c.part.rels.values() if 'hyperlink' in r.reltype)
    print(f"OK {dst} | parrafos={len(c.paragraphs)} tablas={len(c.tables)} links={links} palabras~{words}")

if __name__ == '__main__':
    convert(sys.argv[1], sys.argv[2])
