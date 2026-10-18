# -*- coding: utf-8 -*-
"""
Genera 3 diapositivas de "Preguntas frecuentes" (anexo) usando el mismo
formato del template "Slides template - OWASP UY.pptx"
(layout CUSTOM_1_2 'Title and body': titulo 32pt navy, barra azul,
cuerpo con vinetas • 26pt Lato navy, pie con fuente y numeracion).

Reutiliza los helpers de make_slide_owasp_uy.py.

Salida:
    diapositivas_faq_owasp_uy.pptx   -> 3 slides "FAQ 1 / 3" ... "FAQ 3 / 3"
"""
import sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.text import PP_ALIGN

from make_slide_owasp_uy import (
    TEMPLATE, NAVY, GRAY, get_ph, fill_body, delete_slide,
)


def add_footnote(slide, source, label):
    left, top, width, height = Emu(571550), Emu(5830000), Emu(8250900), Emu(300000)
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Fuente: " + source
    r.font.size = Pt(9)
    r.font.name = "Lato"
    r.font.color.rgb = GRAY

    nb = slide.shapes.add_textbox(Emu(10600000), Emu(5830000), Emu(1100000), Emu(300000))
    ntf = nb.text_frame
    np_ = ntf.paragraphs[0]
    np_.alignment = PP_ALIGN.RIGHT
    nr = np_.add_run()
    nr.text = label
    nr.font.size = Pt(9)
    nr.font.name = "Lato"
    nr.font.color.rgb = GRAY


def build(specs, out_path):
    prs = Presentation(TEMPLATE)
    lay = prs.slide_layouts[3]            # CUSTOM_1_2 = "Title and body"

    for spec in specs:
        slide = prs.slides.add_slide(lay)

        title_ph = get_ph(slide, idx=1)
        ttf = title_ph.text_frame
        ttf.word_wrap = True
        ttf.paragraphs[0].text = spec["title"]

        body_ph = get_ph(slide, idx=2)
        fill_body(body_ph, spec["bullets"])

        add_footnote(slide, spec["source"], spec["label"])
        slide.notes_slide.notes_text_frame.text = spec["guion"]

    # borrar las 4 slides de muestra del template
    for _ in range(4):
        delete_slide(prs, 0)

    prs.save(out_path)
    print("OK ->", out_path, "| slides:", len(prs.slides._sldIdLst))
    return prs


FAQ = [
    {
        "label": "FAQ 1 / 3",
        "title": "Preguntas frecuentes (1/3): formatos y cifrado",
        "bullets": [
            ("Base64 no es cifrado — ",
             "es una codificación sin clave; el cifrado real de JOSE es JWE, y un JWT común es JWS: firmado, no sellado."),
            ("HMAC-SHA1 — ",
             "MAC simétrico (clave + hash, RFC 2104); OAuth 1.0 firmaba cada request con él; hoy se prefiere SHA-256."),
            ("JOSE — ",
             "JWS (firma), JWE (cifrado), JWK (claves), JWA (algoritmos) y JWT (el token con claims)."),
            ("Claims, amr y acr — ",
             "amr (RFC 8176) dice con qué métodos; acr, con qué garantía. La API no sabe si hubo MFA salvo que los revise."),
        ],
        "source": "RFC 2104; RFC 7515–7519; RFC 8176 (2017); RFC 8725 (2020).",
        "guion": (
            "Tres aclaraciones de formato que siempre aparecen. Uno: base64 no "
            "es cifrado, es una codificación sin clave; en JOSE el cifrado real "
            "es JWE, y un JWT común es JWS, o sea firmado pero legible. Dos: "
            "HMAC-SHA1 es un MAC simétrico, clave más hash, que OAuth 1.0 usaba "
            "para firmar cada request; hoy se prefiere SHA-256. Tres: el payload "
            "de un JWT son claims, y amr dice con qué métodos se autenticó "
            "mientras acr dice con qué garantía. La API no sabe que hubo MFA "
            "salvo que el emisor los ponga y el validador los revise."
        ),
    },
    {
        "label": "FAQ 2 / 3",
        "title": "Preguntas frecuentes (2/3): ataques y defensas",
        "bullets": [
            ("JWT mal validado — ",
             "alg:none, confusión HS/RS, inyección de kid. La política de verificación sale de la config, no del token."),
            ("Prueba de posesión — ",
             "mTLS atado al certificado (RFC 8705) o DPoP (RFC 9449): robar el token no alcanza sin la clave privada."),
            ("Passkeys — ",
             "el origen va atado criptográficamente a la firma; resisten el phishing por diseño (WebAuthn, 2019)."),
            ("AiTM e infostealers — ",
             "roban la sesión ya autenticada o la cookie del navegador y la reutilizan: pass-the-cookie."),
        ],
        "source": "RFC 8725 (2020); RFC 8705 (2020); RFC 9449 (2023); WebAuthn (2019).",
        "guion": (
            "Cuatro respuestas sobre ataque y defensa. Un JWT mal validado no es "
            "culpa del token sino del validador: alg:none, confusión de RS256 a "
            "HS256, inyección de kid; la regla es que la política de verificación "
            "sale de la configuración. La defensa de fondo es exigir prueba de "
            "posesión: mTLS con el token atado al certificado, o DPoP firmando "
            "cada request, así robar el token no alcanza. Las passkeys resisten "
            "el phishing porque el origen va atado criptográficamente a la "
            "firma. Y lo que no se cubre es el robo del artefacto ya autenticado: "
            "AiTM e infostealers, el pass-the-cookie."
        ),
    },
    {
        "label": "FAQ 3 / 3",
        "title": "Preguntas frecuentes (3/3): OAuth, HTTP y autorización",
        "bullets": [
            ("¿El bearer es exclusivo de OAuth? — ",
             "no: el esquema Bearer es del marco HTTP; OAuth puede emitir tokens sender-constrained, no bearer puro."),
            ("Flujos OAuth 2 — ",
             "todos emiten access token; Implicit y Password (ROPC) quedaron retirados por RFC 9700 (2025)."),
            ("Versiones de HTTP — ",
             "no crean el problema del bearer, pero cambian la superficie: compresión de cabeceras, inspección binaria, smuggling."),
            ("BOLA / BFLA / BOPLA — ",
             "autorización a nivel de objeto, función y propiedad. Autenticar no es autorizar."),
        ],
        "source": "RFC 6749 / RFC 6750; RFC 9700 (2025); OWASP API Top 10 (2023).",
        "guion": (
            "Tres cierres. El bearer no es exclusivo de OAuth: el esquema vive "
            "en el marco de autenticación de HTTP y cualquier API puede usarlo; "
            "y OAuth puede emitir tokens sender-constrained, que no son bearer "
            "puro. Todos los flujos de OAuth 2 emiten un access token; Implicit "
            "y Password quedaron retirados por RFC 9700 en 2025. Las versiones "
            "de HTTP no crean el problema del bearer, pero cambian la "
            "superficie: compresión de cabeceras, inspección de binario, "
            "smuggling. Y por último, BOLA, BFLA y BOPLA: autorización a nivel "
            "de objeto, función y propiedad. Autenticar no es autorizar."
        ),
    },
]


def numbered(total, start):
    """Copia del FAQ numerada como continuacion del mazo (start .. total)."""
    import copy
    out = copy.deepcopy(FAQ)
    for i, s in enumerate(out):
        s["label"] = "%d / %d" % (start + i, total)
    return out


if __name__ == "__main__":
    import io
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

    # Anexo independiente
    build(FAQ, r"diapositivas_faq_owasp_uy.pptx")
    # Como continuacion del mazo de 28 -> slides 29, 30 y 31 de 31
    build(numbered(31, 29), r"diapositivas_faq_en_28.pptx")
    # Como continuacion del mazo de 25 -> slides 26, 27 y 28 de 28
    build(numbered(28, 26), r"diapositivas_faq_en_25.pptx")

    print("  Nota: el guion de cada slide esta en las notas del ponente.")
